"""Build an editable laboratory report from verified evidence."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "output" / "evidence"


def main():
    results = json.loads((EVIDENCE / "verification.json").read_text(encoding="utf-8"))
    gui = json.loads((EVIDENCE / "gui-smoke.json").read_text(encoding="utf-8"))
    doc = Document()
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    markdown = []

    def heading(text, level=1):
        doc.add_heading(text, level=level)
        markdown.append("#" * (level + 1) + " " + text + "\n")

    def paragraph(text):
        doc.add_paragraph(text)
        markdown.append(text + "\n")

    def table(headers, rows, widths):
        target = doc.add_table(rows=1, cols=len(headers))
        target.autofit = False
        for column, width in zip(target.columns, widths):
            column.width = Cm(width)
        for index, value in enumerate(headers):
            target.rows[0].cells[index].text = str(value)
        for row in rows:
            cells = target.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        for row_index, row in enumerate(target.rows):
            props = row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            props.append(no_split)
            if row_index == 0:
                repeat = OxmlElement("w:tblHeader")
                props.append(repeat)
            for column_index, cell in enumerate(row.cells):
                cell.width = Cm(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                properties.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), "80")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E9EEF2")
                    properties.append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.2
                    for r in p.runs:
                        r.font.size = Pt(10)
                        r.bold = row_index == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        markdown.append("| " + " | ".join(headers) + " |")
        markdown.append("| " + " | ".join("---" for _ in headers) + " |")
        markdown.extend("| " + " | ".join(str(v) for v in row) + " |" for row in rows)
        markdown.append("")

    def new_page():
        doc.add_page_break()
        markdown.append("\n")

    def figure(filename, caption):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(EVIDENCE / filename), width=Cm(15.2))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(10)
        markdown.append(f"![{caption}](evidence/{filename})\n")

    title = "密码学算法应用中级实验报告"
    doc.add_paragraph(title, "Title")
    markdown.append("# " + title + "\n")
    paragraph("实验项目：密语安全通信系统")
    paragraph("姓名：____________    学号：____________    班级：____________")
    paragraph("实验日期：2026 年 10 月 16 日")
    heading("一 实验目标与完成范围")
    paragraph("本实验设计并实现一个带图形界面的双端安全通信程序，综合应用对称加密、非对称加密、Hash 和密钥交换算法，完成文字、图片及视频文件的双向传输。程序优先采用 SM2、SM3 和 SM4，临时密钥协商采用 X25519。双方在身份核验和密钥确认完成后才能发送应用数据。")
    paragraph("实验重点是说明密码算法如何在通信流程中协作：SM2 负责认证握手并加密随机挑战，X25519 和 HKDF-SM3 生成会话密钥，SM4 加密应用内容，HMAC-SM3 校验传输记录，SM3 检查文件整体一致性。")
    table(["中级实验要求", "对应实现"], [
        ("可视化交互", "中文 Tkinter 界面，通信、会话身份、安全实验三个页面"),
        ("对称加密", "SM4-CBC 与随机 IV，使用 PKCS7 填充"),
        ("非对称加密", "SM2 加密双方 32 字节随机挑战并验证解密回显"),
        ("Hash 与消息认证", "SM3 文件摘要与 HMAC-SM3 逐记录认证"),
        ("密钥交换", "X25519 临时密钥协商与 HKDF-SM3 密钥派生"),
        ("文字 图片 视频", "双向中文消息、图片预览、MP4 文件加密收发"),
        ("复现与原理验收", "源代码、样本、自动测试、实际截图和验收问答"),
    ], [4.0, 13.0])
    heading("实验环境", 2)
    paragraph(f"实测系统为 Windows 11，Python {results['python']}，cryptography {results['cryptography']}，gmssl 3.2.2，Pillow 12.3.0。界面使用 Python 自带 Tkinter，通信使用标准 socket；测试采用 unittest。两台端点在本机回环地址上完成本次复现，局域网模式由相同 TCP 协议支持，本报告不将回环测试结果等同于两台物理电脑的实测结果。")

    new_page()
    heading("二 密码方案与通信流程")
    table(["阶段", "算法与用途", "关键数据"], [
        ("身份认证", "SM2 签名与 SM3", "签名绑定角色、身份、临时公钥和随机数"),
        ("临时协商", "X25519", "双方仅交换临时公钥，私钥保留在本地"),
        ("密钥派生", "HKDF-SM3", "共享秘密与完整握手摘要作为输入"),
        ("公钥加密", "SM2 加密随机挑战", "对端私钥解密并回显，双方各执行一次"),
        ("数据加密", "SM4-CBC", "每条记录生成新的 16 字节随机 IV"),
        ("记录认证", "HMAC-SM3", "认证会话标识、连续序号、IV 和密文"),
        ("本地私钥保护", "scrypt 与 AES-GCM", "口令派生密钥，私钥以认证密文保存"),
    ], [2.6, 5.3, 9.1])
    heading("握手过程", 2)
    for text in [
        "1. 双方解锁本地长期 SM2 身份，每次连接独立生成 X25519 临时密钥和 32 字节随机数。",
        "2. 构造带版本、角色、名称、身份公钥、临时公钥和随机数的握手消息，对规范化内容进行 SM2 签名后交换。",
        "3. 验证对端字段、角色、曲线点和 SM2 签名；首次连接通过可信渠道核对完整 SM3 身份指纹，保存信任记录。已保存的同名身份发生改变时拒绝连接。",
        "4. 用本地 X25519 私钥和对端临时公钥计算共享秘密。HKDF-SM3 输出 96 字节，划分为两个方向各 16 字节 SM4 密钥和 32 字节 HMAC 密钥。",
        "5. 双方分别产生新的 32 字节随机挑战，用对端 SM2 公钥加密后放入经过 SM4 加密和 HMAC 认证的记录中发送。",
        "6. 接收方使用本地 SM2 私钥解密挑战并验证 C3 摘要，通过加密通道回显明文；两端检查回显与原挑战一致后，界面显示安全连接。",
    ]:
        paragraph(text)
    paragraph("SM2 公钥加密体现非对称加密的作用，SM2 签名认证协商参数，两者职责不同。批量媒体数据使用 SM4，避免对视频逐字节执行高成本的非对称运算。")

    new_page()
    heading("三 实现方法与关键代码")
    heading("加密记录和重放防护", 2)
    paragraph("Channel.seal 对明文执行 PKCS7 填充，以随机 IV 进行 SM4-CBC 加密，再计算 HMAC-SM3。密文记录依次包含 8 字节序号、16 字节 IV、密文和 32 字节认证码。认证计算还包含会话标识，因此旧会话记录不能迁移到新会话。")
    paragraph("Channel.open 先检查长度，再用恒定时间比较函数验证认证码，然后检查序号是否为下一条预期记录。全部验证成功后才解密和去填充；失败不会推进接收序号。发送和接收使用独立密钥及独立序号。")
    heading("文件分块与接收发布", 2)
    paragraph("send_file 发送经过认证的文件名、大小、种类和传输标识，随后按 64 KiB 分块读取和传输，同时增量计算 SM3。接收端将数据写入随机名称的临时文件，并限制累计长度。结束消息中的 SM3 摘要经加密认证后传送。")
    paragraph("_file_end 仅在实际字节数和 SM3 都正确时发布文件，并发送完成回执；发送方收到回执后才显示对端已校验接收。同名文件追加编号，避免覆盖；危险路径名和设备名被拒绝。连接中断或校验失败时清理未完成文件。")
    heading("界面线程与密钥管理", 2)
    paragraph("socket 收发及文件传输在工作线程执行，界面主线程通过事件队列更新聊天记录和进度，避免长文件传输冻结窗口。首次指纹核验使用非阻塞对话框，双方可同时查看信息。文字草稿在发送成功事件到达后才清除；失败时保留供重试。")
    paragraph("身份私钥以 JSON 容器保存。scrypt 使用 n=32768、r=8、p=1 派生 32 字节保护密钥，AES-GCM 使用随机盐与随机 nonce 保存私钥密文，并将公钥作为附加认证数据。口令不写盘，会话密钥不写入日志。")
    heading("库适配与模块划分", 2)
    paragraph("gmssl 3.2.2 的默认随机源不适合密码用途，因此适配层统一使用 secrets；同时补上 SM2 曲线点验证、解密结果 C3 比较，并保留可能以 0 或 4 开头的原始公钥坐标。SM2 曲线运算与加解密核心仍调用库完成。")
    paragraph("模块分工：crypto.py 实现身份与密码操作；protocol.py 负责帧、会话和文件传输；gui.py 与 main.py 负责界面和入口；tests 与 tools 提供异常路径测试、样本和证据生成。")

    new_page()
    heading("四 测试方法与实测结果")
    paragraph(f"本次 unittest 运行共 {results['tests_passed']} 项测试，全部通过。原始输出保存为 output/evidence/tests-output.txt，环境与性能记录保存为 verification.json，图形界面综合测试结果保存为 gui-smoke.json。测试既验证正确收发，也检查无效输入是否被拒绝。")
    table(["测试组", "样本或操作", "结果"], [
        ("标准摘要", "SM3 对 abc 的标准向量", "一致"),
        ("SM2 签名与加密", "消息变更、错误密钥、C3 变更、非法点", "拒绝"),
        ("私钥文件", "正确口令与错误口令", "正确解锁或拒绝"),
        ("握手", "替换临时公钥、反射角色、错误挑战回显", "拒绝"),
        ("加密记录", "双向收发、篡改、重放、跨会话、乱序与截断", "符合预期"),
        ("传输边界", "TCP 分片、超限帧、危险文件名、身份改变", "正确处理"),
        ("文件与异常", "多块二进制、空文件、同名文件、摘要错误、中断", "一致或拒绝并清理"),
        ("界面综合", "中文聊天、PNG、MP4、预览、窄窗口、失败草稿", "通过"),
    ], [3.0, 10.4, 3.6])
    paragraph(f"真实 PNG 样本为 {gui['image']['bytes']:,} 字节，真实 H.264 MP4 视频为 {gui['video']['bytes']:,} 字节、时长 4 秒。界面测试通过实际 TCP 传输后逐字节比较源文件与接收文件，两类媒体均一致，图片预览正常。视频按文件保存，由操作系统播放器播放。")
    heading("效率测量", 2)
    table(["指标", "本次回环测量"], [
        ("握手与双向确认", f"{results['handshake_seconds']:.4f} 秒"),
        ("文件大小", f"{results['file_bytes']:,} 字节 即 8 MiB"),
        ("传输至校验回执耗时", f"{results['file_seconds']:.4f} 秒"),
        ("有效载荷吞吐量", f"{results['throughput_MiB_s']:.2f} MiB/s"),
        ("接收文件逐字节比较", "完全一致"),
    ], [7.0, 10.0])
    paragraph("该测量仅为当前电脑回环地址上的一次运行，包含分块加密、接收写盘、摘要检查和完成回执，不代表局域网或公网的稳定吞吐量。分块方式使传输内存随块大小增长，而不是随整个视频大小增长；程序不提供断点续传或实时视频通话。")

    new_page()
    heading("五 文字与媒体通信结果")
    paragraph("图 1 展示 Bob 接收到 Alice 的加密消息、PNG 图片和 MP4 视频文件后的界面。图片显示缩略预览，媒体记录同时显示接收文件大小和 SM3 摘要。对端界面在收到校验回执后显示已校验接收。")
    figure("bob-chat.png", "图 1 实际双端通信中的接收界面")
    paragraph("测试使用不含个人信息的样本，接收端最终文件与发送端逐字节一致。图中身份指纹属于本次临时测试身份，与正式启动时由用户口令保护的身份相互独立。")

    new_page()
    heading("六 篡改与重放检测结果")
    paragraph("安全实验为独立临时会话：先修改密文中的一个比特，确认 HMAC 检查拒绝；再将已成功接收的记录重新提交，确认严格序号检查拒绝。图 2 展示两种拒绝结果和不含明文内容的密文记录样本。")
    figure("security-lab.png", "图 2 实际运行的密文篡改与重放实验")
    paragraph("实验按钮不会破坏正在使用的真实聊天会话。独立测试展示两种失败路径，自动化测试进一步检查跨会话记录、无效身份、错误文件摘要和连接中断；这些检查共同支持程序的实现结论，但不能替代安全协议的独立审计。")

    new_page()
    heading("七 复现步骤与分析")
    for text in [
        "1. 安装含 Tcl/Tk 的 Python 3.12，运行 install.bat 安装依赖；当前实验目录已准备好独立环境。",
        "2. 运行 start_demo.bat，设置或输入身份口令，打开 Alice 与 Bob 两个窗口，程序自动选择本机端口连接。",
        "3. 首次连接通过可信渠道核对双方完整身份指纹并分别确认；已保存的身份在后续连接自动核对。",
        "4. 双向发送文字，再发送 samples 中的 PNG 和 MP4，检查预览、摘要、完成提示及接收目录。",
        "5. 打开会话身份页检查算法用途，再运行安全实验并导出日志；断开重连后会话标识应改变。",
        "6. 运行 run_tests.bat，或 tools/verify_lab.py 重新生成测试输出和回环测量。两台电脑使用时由 Bob 监听，Alice 填写 Bob 的真实局域网地址连接。",
    ]:
        paragraph(text)
    heading("安全性和工程边界", 2)
    paragraph("机密性来自 SM4 加密，消息来源和完整性来自 HMAC-SM3；身份可信性依赖 SM2 签名、已确认指纹及首次可信核验。若首次核验被省略，仅点击信任不能排除中间人攻击。普通 SM3 摘要不能独立提供来源认证。")
    paragraph("每次连接重新生成临时协商密钥，并把双方完整握手绑定到派生过程，避免复用旧会话。程序在断开后释放密钥对象引用，但 Python 不保证物理内存清零，因此不宣称实现了可靠内存擦除。接收文件在本地以正常明文保存，传输加密不等于文件静态加密。")
    paragraph("程序限制单条文字 32 KiB、文件 2 GiB、最大传输帧和同时接收数量；同名文件不会覆盖，未通过摘要校验的文件不会发布。强制结束进程仍可能留下未完成临时文件；磁盘容量、系统访问权限和网络条件会影响传输。本应用属于教学自定义协议，未经过商用密码认证或第三方安全审计。")
    heading("实验结论", 2)
    paragraph(f"本实验完成了图形化双端安全通信，覆盖文字、图片和视频文件，以及中级要求中的四类密码算法。本次 {results['tests_passed']} 项自动测试和实际界面复现通过，展示了正确解密、文件一致性及异常输入拒绝。程序把身份认证、临时密钥协商、数据机密性与记录认证分开实现，便于结合关键代码解释各算法的职责。")
    heading("参考资料", 2)
    paragraph("[1] 吕 密码学算法应用实验验收标准与报告评分细则，第 1 至 3 页。\n[2] cryptography 项目文档 https://cryptography.io/\n[3] gmssl 3.2.2 项目与源代码 https://pypi.org/project/gmssl/3.2.2/\n[4] RFC 7748 Elliptic Curves for Security；RFC 5869 HKDF。")
    output = ROOT / "output"
    path = output / "密码学中级实验报告.docx"
    doc.core_properties.title = title
    doc.core_properties.subject = "SM2 SM3 SM4 X25519 安全通信实验"
    doc.core_properties.author = ""
    doc.save(path)
    (output / "密码学中级实验报告.md").write_text("\n".join(markdown), encoding="utf-8")
    print(path)


if __name__ == "__main__":
    main()
